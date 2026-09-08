import openpyxl
import json
from collections import defaultdict
from datetime import datetime

def safe_float(v):
    try: return float(v) if v not in (None, 'None', '') else 0.0
    except: return 0.0

def safe_date(v):
    if v is None or v == 'None': return None
    if isinstance(v, datetime): return v
    try:
        s = str(v).strip()
        for fmt in ('%Y-%m-%d %H:%M:%S.%f', '%Y-%m-%d %H:%M:%S', '%Y-%m-%d'):
            try: return datetime.strptime(s[:19], fmt)
            except: pass
    except: pass
    return None

base = "C:/Users/Montolivo/Dropbox/Asesorias PYMES/1. Proyectos Actuales/2. MONTOLIVO/2. Proyecto/Indicadores Montolivo/Compras/Bases para actualizar Claude/"

months_all = ["2025-01","2025-02","2025-03","2025-04","2025-05","2025-06","2025-07","2025-08",
              "2025-09","2025-10","2025-11","2025-12","2026-01","2026-02","2026-03","2026-04","2026-05","2026-06","2026-07","2026-08","2026-09"]
month_idx = {m:i for i,m in enumerate(months_all)}
last4 = ["2026-06","2026-07","2026-08","2026-09"]

# Unidades base "limpias": una unidad de compra que siempre equivale a un valor fijo en la unidad
# base del ítem (GR o UNID), sin importar el ítem específico.
CLEAN_BASE = {
    'GR': ('GR', 1.0), 'KG': ('GR', 1000.0), 'LB': ('GR', 453.592),
    'UNID': ('UNID', 1.0), 'UND': ('UNID', 1.0),
}

# ════════════════════════════════════════════════════════════════
# 0) Homologación oficial por Referencia — hoja "Und de medida" (misma fuente
#    que usan Alex y JR): Referencia -> (U.M. inventario, factor, U.M. de orden)
# ════════════════════════════════════════════════════════════════
print("Leyendo hoja 'Und de medida' (homologación oficial por Referencia)...")
wb1 = openpyxl.load_workbook(base + "Histórico de compras por item.xlsx", data_only=True, read_only=True)
ws_um = wb1['Und de medida']
um_lookup = {}
for i, r in enumerate(ws_um.iter_rows(values_only=True)):
    if i == 0:
        continue
    ref = str(r[0]).strip() if r[0] else ''
    if not ref or ref in um_lookup:
        continue
    um_invent = (r[2] or '').strip()
    factor = r[3]
    um_orden = str(r[4]).strip() if r[4] is not None else ''
    if isinstance(factor, (int, float)) and factor > 0 and um_invent in ('GR', 'UNID'):
        um_lookup[ref] = (um_invent, float(factor), um_orden)
print(f"  Ítems con homologación oficial: {len(um_lookup)}")

# ════════════════════════════════════════════════════════════════
# 1) PRECIO UNITARIO MENSUAL POR (PROVEEDOR, ITEM) — ultimos 4 meses
#    Precio = valor bruto / cantidad HOMOLOGADA (kg o unidad base), no cantidad cruda.
# ════════════════════════════════════════════════════════════════
print("Leyendo historico de compras (precio unitario homologado por proveedor+item+mes)...")
ws2 = wb1.active

pair_total_recent   = defaultdict(float)                       # TODO el bruto (homologable o no) ultimos 4 meses -> "cantidad comprada en $"
raw_homolog_rows    = defaultdict(list)                        # key -> [(mes, bruto, qty_base, base_unit), ...]

headers2 = None
n = 0
n_homolog = n_no_homolog = 0
for i, row in enumerate(ws2.iter_rows(values_only=True)):
    if i == 0:
        headers2 = [str(c).strip() if c else '' for c in row]
        continue
    r = dict(zip(headers2, row))
    fecha = safe_date(r.get('Fecha'))
    if not fecha: continue
    key_month = f"{fecha.year}-{fecha.month:02d}"
    if key_month not in last4: continue
    prov = str(r.get('Razon social proveedor') or r.get('Razón social proveedor') or '').strip()
    item = str(r.get('Desc. item') or '').strip()
    ref = str(r.get('Referencia') or '').strip()
    cant = safe_float(r.get('Cantidad'))
    bruto = safe_float(r.get('Valor bruto local'))
    um = str(r.get('U.M.') or '').strip()
    if not prov or not item or cant == 0: continue
    n += 1
    key = (prov, item)
    pair_total_recent[key] += bruto

    info = um_lookup.get(ref)
    qty_base = None
    base_unit = None
    if info:
        um_invent, factor_orden, um_orden = info
        base_unit = um_invent
        if um == um_orden:
            qty_base = cant * factor_orden
        elif um in CLEAN_BASE and CLEAN_BASE[um][0] == um_invent:
            qty_base = cant * CLEAN_BASE[um][1]
    else:
        if um in CLEAN_BASE:
            base_unit = CLEAN_BASE[um][0]
            qty_base = cant * CLEAN_BASE[um][1]

    if qty_base and qty_base > 0:
        raw_homolog_rows[key].append((key_month, bruto, qty_base, base_unit))
        n_homolog += 1
    else:
        n_no_homolog += 1

wb1.close()
print(f"  Filas ultimos 4 meses: {n}  Homologadas: {n_homolog}  Sin homologar: {n_no_homolog}")
print(f"  Pares proveedor+item con datos homologados: {len(raw_homolog_rows)}")

# ── Guardia contra errores de digitación puntuales ──
# Un solo registro con una Cantidad mal tecleada (ej. "15000" en vez de "15") puede disparar la
# cantidad homologada de ese mes miles de veces por encima de lo real, hundiendo el precio
# calculado a casi $0. Se descarta una fila si su precio implícito (bruto/qty) se desvía más de
# 20x de la mediana de precios del propio par en el período — margen amplio para no censurar
# variación de precio real (ej. verduras con oferta/escasez), solo errores evidentes.
OUTLIER_FACTOR = 20
pair_month_bruto_h = defaultdict(lambda: defaultdict(float))
pair_month_cant_h  = defaultdict(lambda: defaultdict(float))
pair_base_unit     = {}
n_outliers = 0
for key, filas in raw_homolog_rows.items():
    precios_fila = sorted(b / q for _, b, q, _ in filas if q > 0)
    if not precios_fila:
        continue
    mediana = precios_fila[len(precios_fila)//2]
    for key_month, bruto, qty_base, base_unit in filas:
        p = bruto / qty_base
        if mediana > 0 and (p < mediana / OUTLIER_FACTOR or p > mediana * OUTLIER_FACTOR):
            n_outliers += 1
            continue
        pair_month_bruto_h[key][key_month] += bruto
        pair_month_cant_h[key][key_month] += qty_base
        pair_base_unit[key] = base_unit
print(f"  Filas descartadas por outlier extremo (posible error de digitación): {n_outliers}")

# Construir lista de pares con precio unitario homologado por mes, exigiendo al menos 2 meses con datos
pairs_out = []
for key, monthmap in pair_month_bruto_h.items():
    prov, item = key
    months_with_data = [m for m in last4 if monthmap.get(m,0) > 0 and pair_month_cant_h[key].get(m,0) > 0]
    if len(months_with_data) < 2:
        continue
    base_unit = pair_base_unit[key]
    # Para items de peso mostramos $/kg (mas legible); para items de unidad, $/unidad.
    divisor = 1000.0 if base_unit == 'GR' else 1.0
    unidad_mostrada = 'kg' if base_unit == 'GR' else 'unid'
    precios = {}
    for m in last4:
        c = pair_month_cant_h[key].get(m, 0)
        b = monthmap.get(m, 0)
        precios[m] = round(b / (c / divisor), 2) if c > 0 else None
    # Variación = mes actual (último de la ventana) vs. promedio de los 3 meses anteriores
    # (se promedian solo los meses anteriores que sí tengan precio disponible).
    mes_actual = precios[last4[-1]]
    anteriores = [precios[m] for m in last4[:-1] if precios[m] is not None]
    promedio_anterior = sum(anteriores) / len(anteriores) if anteriores else None
    variacion_pct = round((mes_actual - promedio_anterior) / promedio_anterior * 100, 1) if (mes_actual is not None and promedio_anterior) else None
    pairs_out.append({
        "proveedor": prov, "item": item, "um": unidad_mostrada,
        "precios": [precios[m] for m in last4],
        "variacionPct": variacion_pct,
        "totalRecienteBruto": round(pair_total_recent[key], 0)
    })

pairs_out.sort(key=lambda x: -x["totalRecienteBruto"])
print(f"  Pares con >=2 meses de datos homologados: {len(pairs_out)}")

# ════════════════════════════════════════════════════════════════
# 2) ORDENES DE COMPRA: conteo y valor mensual (2025 vs 2026)
# ════════════════════════════════════════════════════════════════
print("Leyendo ordenes de compra (conteo mensual de OC unicas)...")
wb3 = openpyxl.load_workbook(base + "Orden de compra por item.xlsx", data_only=True, read_only=True)
ws3 = wb3.active

oc_unique_by_month = defaultdict(set)
oc_valor_by_month = defaultdict(float)

headers3 = None
for i, row in enumerate(ws3.iter_rows(values_only=True)):
    if i == 0:
        headers3 = [str(c).strip() if c else '' for c in row]
        continue
    r = dict(zip(headers3, row))
    f_orden = safe_date(r.get('Fecha orden'))
    nro = r.get('Nro orden')
    valor = safe_float(r.get('Valor neto'))
    if f_orden:
        key = f"{f_orden.year}-{f_orden.month:02d}"
        if key in month_idx:
            if nro is not None:
                oc_unique_by_month[key].add(nro)
            oc_valor_by_month[key] += valor

wb3.close()
oc_count_by_month = {m: len(oc_unique_by_month.get(m, set())) for m in months_all}
print("  OC unicas por mes:", oc_count_by_month)

out = {
    "last4Months": last4,
    "priceVariation": pairs_out,
    "ocCountByMonth": oc_count_by_month
}

with open(base + "filter_data_v2.json", "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, separators=(',',':'))

print("Guardado filter_data_v2.json")
