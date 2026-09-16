import openpyxl
import json
from collections import defaultdict
from datetime import datetime

def safe_float(v):
    try: return float(v) if v not in (None, 'None', '') else 0.0
    except: return 0.0

def safe_date(v):
    if v is None or v == 'None': return None
    if isinstance(v, datetime): return v
    try:
        s = str(v).strip()
        for fmt in ('%Y-%m-%d %H:%M:%S.%f', '%Y-%m-%d %H:%M:%S', '%Y-%m-%d'):
            try: return datetime.strptime(s[:19], fmt)
            except: pass
    except: pass
    return None

base = "C:/Users/Montolivo/Dropbox/Asesorias PYMES/1. Proyectos Actuales/2. MONTOLIVO/2. Proyecto/Indicadores Montolivo/Compras/Bases para actualizar Claude/"

months_all = ["2025-01","2025-02","2025-03","2025-04","2025-05","2025-06","2025-07","2025-08",
              "2025-09","2025-10","2025-11","2025-12","2026-01","2026-02","2026-03","2026-04","2026-05","2026-06","2026-07","2026-08","2026-09"]
month_idx = {m: i for i, m in enumerate(months_all)}
NM = len(months_all)

print("Leyendo historico de compras (cruce proveedor x item x mes)...")
wb = openpyxl.load_workbook(base + "Histórico de compras por item.xlsx", data_only=True, read_only=True)
ws = wb[wb.sheetnames[0]]

items_by_prov = defaultdict(lambda: defaultdict(lambda: [0.0]*NM))
provs_by_item = defaultdict(lambda: defaultdict(lambda: [0.0]*NM))

headers = None
n = 0
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i == 0:
        headers = [str(c).strip() if c else '' for c in row]
        continue
    r = dict(zip(headers, row))
    n += 1
    prov = str(r.get('Razon social proveedor') or r.get('Razón social proveedor') or '').strip()
    item = str(r.get('Desc. item') or '').strip()
    if not prov or not item:
        continue
    fecha = safe_date(r.get('Fecha'))
    if not fecha:
        continue
    key = f"{fecha.year}-{fecha.month:02d}"
    idx = month_idx.get(key)
    if idx is None:
        continue
    neto = safe_float(r.get('Valor neto local'))
    items_by_prov[prov][item][idx] += neto
    provs_by_item[item][prov][idx] += neto

wb.close()
print(f"  Filas: {n}  Proveedores: {len(items_by_prov)}  Items: {len(provs_by_item)}")

M = 1_000_000

def round_arr(arr):
    return [round(v/M, 1) for v in arr]

out = {
    "months": months_all,
    "itemsByProv": {p: {it: round_arr(arr) for it, arr in items.items()} for p, items in items_by_prov.items()},
    "provsByItem": {it: {p: round_arr(arr) for p, arr in provs.items()} for it, provs in provs_by_item.items()},
}

with open(base + "extract/prov_item_cross.json", "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, separators=(',', ':'))

print("Guardado prov_item_cross.json")
npairs = sum(len(v) for v in out["itemsByProv"].values())
print("Pares proveedor-item:", npairs)
