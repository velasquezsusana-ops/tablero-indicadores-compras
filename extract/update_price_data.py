import openpyxl, json
from collections import defaultdict
from datetime import datetime
from pyxlsb import open_workbook

def safe_float(v):
    try: return float(v) if v not in (None,'None','') else 0.0
    except: return 0.0

def safe_date(v):
    if v is None or v == 'None': return None
    if isinstance(v, datetime): return v
    try:
        s = str(v).strip()
        for fmt in ('%Y-%m-%d %H:%M:%S.%f','%Y-%m-%d %H:%M:%S','%Y-%m-%d'):
            try: return datetime.strptime(s[:19], fmt)
            except: pass
    except: pass
    return None

BASE = "C:/Users/Montolivo/Dropbox/Asesorias PYMES/1. Proyectos Actuales/2. MONTOLIVO/2. Proyecto/Indicadores Montolivo/Compras/Bases para actualizar Claude/"
HTML_PATH = BASE + "Tablero_Indicadores_Compras.html"

last4 = ["2026-06","2026-07","2026-08","2026-09"]
CLEAN_BASE = {'GR':('GR',1.0),'KG':('GR',1000.0),'LB':('GR',453.592),'UNID':('UNID',1.0),'UND':('UNID',1.0)}

# 1) Factores de conversion -> precios unitarios.xlsb / 'Factores conversión'
print("Leyendo 'Factores conversión'...")
factor_lookup = {}
with open_workbook(BASE + "precios unitarios.xlsb") as wbx:
    with wbx.get_sheet('Factores conversión') as sh:
        for i, row in enumerate(sh.rows()):
            if i == 0: continue
            vals = [c.v for c in row]
            if not vals or not vals[0]: continue
            ref, um_compra, factor, um_base, origen = vals[1], vals[3], vals[4], vals[5], vals[6]
            if origen == 'PENDIENTE' or not ref or not factor: continue
            factor_lookup[(str(ref).strip(), str(um_compra).strip())] = (float(factor), str(um_base).strip())
# Correcciones confirmadas por Compras cuando el factor del Excel está errado (tienen prioridad).
# MP0127 QUESO CREMA: en "Factores conversión" dice 1 CAJA = 1 GR; la caja es de 4.000 g (confirmado oct-2026).
FACTOR_FIX = {('MP0127', 'CAJA'): (4000.0, 'GR')}
factor_lookup.update(FACTOR_FIX)
print(f"  Factores: {len(factor_lookup)}")

# 2) Precio homologado por (Proveedor,Item) - ultimos 4 meses -> 'Histórico de compras por item.xlsx'
print("Recalculando precio homologado ultimos 4 meses...")
wb1 = openpyxl.load_workbook(BASE + "Histórico de compras por item.xlsx", data_only=True, read_only=True)
ws2 = wb1[wb1.sheetnames[0]]

pair_total_recent = defaultdict(float)
pair_ref = {}
raw_homolog_rows = defaultdict(list)
headers2 = None
for i, row in enumerate(ws2.iter_rows(values_only=True)):
    if i == 0:
        headers2 = [str(c).strip() if c else '' for c in row]; continue
    r = dict(zip(headers2, row))
    fecha = safe_date(r.get('Fecha'))
    if not fecha: continue
    key_month = f"{fecha.year}-{fecha.month:02d}"
    if key_month not in last4: continue
    prov = str(r.get('Razon social proveedor') or r.get('Razón social proveedor') or '').strip()
    item = str(r.get('Desc. item') or '').strip()
    ref = str(r.get('Referencia') or '').strip()
    cant = safe_float(r.get('Cantidad'))
    bruto = safe_float(r.get('Valor bruto local'))
    um = str(r.get('U.M.') or '').strip()
    if not prov or not item or cant == 0: continue
    key = (prov, item)
    pair_total_recent[key] += bruto
    pair_ref[key] = ref
    info = factor_lookup.get((ref, um))
    qty_base = base_unit = None
    if info:
        factor, um_base = info
        base_unit, qty_base = um_base, cant * factor
    elif um in CLEAN_BASE:
        base_unit, factor = CLEAN_BASE[um]
        qty_base = cant * factor
    if qty_base and qty_base > 0:
        raw_homolog_rows[key].append((key_month, bruto, qty_base, base_unit))
wb1.close()

OUTLIER_FACTOR = 20
pair_month_bruto_h = defaultdict(lambda: defaultdict(float))
pair_month_cant_h = defaultdict(lambda: defaultdict(float))
pair_base_unit = {}
for key, filas in raw_homolog_rows.items():
    precios_fila = sorted(b/q for _,b,q,_ in filas if q > 0)
    if not precios_fila: continue
    mediana = precios_fila[len(precios_fila)//2]
    for key_month, bruto, qty_base, base_unit in filas:
        p = bruto / qty_base
        if mediana > 0 and (p < mediana/OUTLIER_FACTOR or p > mediana*OUTLIER_FACTOR): continue
        pair_month_bruto_h[key][key_month] += bruto
        pair_month_cant_h[key][key_month] += qty_base
        pair_base_unit[key] = base_unit

pairs_out = []
for key, monthmap in pair_month_bruto_h.items():
    prov, item = key
    months_with_data = [m for m in last4 if monthmap.get(m,0) > 0 and pair_month_cant_h[key].get(m,0) > 0]
    if len(months_with_data) < 2: continue
    base_unit = pair_base_unit[key]
    divisor = 1000.0 if base_unit == 'GR' else 1.0
    unidad_mostrada = 'kg' if base_unit == 'GR' else 'unid'
    precios = {}
    for m in last4:
        c = pair_month_cant_h[key].get(m,0); b = monthmap.get(m,0)
        precios[m] = round(b/(c/divisor),2) if c > 0 else None
    mes_actual = precios[last4[-1]]
    anteriores = [precios[m] for m in last4[:-1] if precios[m] is not None]
    prom_ant = sum(anteriores)/len(anteriores) if anteriores else None
    var_pct = round((mes_actual-prom_ant)/prom_ant*100,1) if (mes_actual is not None and prom_ant) else None
    # Cantidad, gasto y sobrecosto/ahorro del MISMO proveedor+item y del MISMO mes actual que las columnas de precio
    # (antes se tomaban de "precios unitarios.xlsb" / AA-AE por Referencia: todos los proveedores y otro mes).
    m_ult = last4[-1]
    cant_act = pair_month_cant_h[key].get(m_ult, 0) / divisor
    gasto_act = monthmap.get(m_ult, 0)
    tiene_act = cant_act > 0 and gasto_act > 0
    gasto_prom = cant_act * prom_ant if (tiene_act and prom_ant) else None
    pairs_out.append({
        "proveedor": prov, "item": item, "referencia": pair_ref.get(key,""), "um": unidad_mostrada,
        "precios": [precios[m] for m in last4], "variacionPct": var_pct,
        "totalRecienteBruto": round(pair_total_recent[key],0),
        "pctVarVsProm3M": (mes_actual-prom_ant)/prom_ant if (mes_actual is not None and prom_ant) else None,
        "cantidadMesActual": round(cant_act,2) if tiene_act else None, "cantidadMesActualUM": unidad_mostrada,
        "gastoRealMesActual": gasto_act if tiene_act else None,
        "gastoPrecioProm3M": gasto_prom,
        "sobrecostoAhorro": (gasto_act - gasto_prom) if gasto_prom is not None else None,
    })
pairs_out.sort(key=lambda x: -x["totalRecienteBruto"])
print(f"  Pares homologados (>=2 meses): {len(pairs_out)}")

mes_actual_label = f"Cantidad mes actual ({last4[-1]})"

# 4) Preservar ocCountByMonth (tabla "Comparativo Anual", no depende de este flujo)
html = open(HTML_PATH, encoding='utf-8').read()
marker = 'let PRICE_DATA = '
start = html.index(marker) + len(marker)
old_obj, end_idx = json.JSONDecoder().raw_decode(html, start)
assert html[end_idx] == ';', "Delimitador ';' no encontrado - abortar sin escribir"
oc_count = old_obj['ocCountByMonth']

new_obj = {"last4Months": last4, "mesActualAAAE": mes_actual_label, "priceVariation": pairs_out, "ocCountByMonth": oc_count}
new_json = json.dumps(new_obj, ensure_ascii=False, separators=(',',':'))
new_html = html[:start] + new_json + html[end_idx:]
open(HTML_PATH, 'w', encoding='utf-8').write(new_html)
print(f"PRICE_DATA actualizado en HTML: {len(pairs_out)} pares proveedor-item.")
