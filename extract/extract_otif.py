import openpyxl
import json
from collections import defaultdict
from datetime import datetime

def safe_float(v):
    try: return float(v) if v not in (None, 'None', '') else 0.0
    except: return 0.0

def safe_date(v):
    if v is None or v == 'None': return None
    if isinstance(v, datetime): return v
    try:
        s = str(v).strip()
        for fmt in ('%Y-%m-%d %H:%M:%S.%f', '%Y-%m-%d %H:%M:%S', '%Y-%m-%d'):
            try: return datetime.strptime(s[:19], fmt)
            except: pass
    except: pass
    return None

base = "C:/Users/Montolivo/Dropbox/Asesorias PYMES/1. Proyectos Actuales/2. MONTOLIVO/2. Proyecto/Indicadores Montolivo/Compras/Bases para actualizar Claude/"

months_all = ["2025-01","2025-02","2025-03","2025-04","2025-05","2025-06","2025-07","2025-08",
              "2025-09","2025-10","2025-11","2025-12","2026-01","2026-02","2026-03","2026-04","2026-05","2026-06","2026-07","2026-08","2026-09"]
month_idx = {m: i for i, m in enumerate(months_all)}
NM = len(months_all)

def zeros():
    return [0.0] * NM

def new_bucket():
    return {"n": zeros(), "ifOk": zeros(), "otOk": zeros(), "otifOk": zeros(),
            "valorEval": zeros(), "valorOtifOk": zeros()}

# ════════════════════════════════════════════════════════════════
# Analisis por LINEA (renglon de "Orden de compra por item.xlsx"), con UN SOLO
# denominador N para OT, IF y OTIF: renglones no Anulados cuya orden (Nro orden)
# tiene al menos una entrega vinculada en el Historico (Docto. orden) de la que se
# puede derivar si llego a tiempo. Asi los tres % son directamente comparables
# (OTIF <= min(OT, IF) siempre) y "OTIF en $" usa el mismo universo de renglones.
#   - OT  = la orden llego a tiempo (TODAS sus entregas vinculadas <= fecha prometida)
#   - IF  = el propio renglon quedo completo (Cant. entrada >= Cant. ordenada)
#   - OTIF = OT y IF a la vez
# Los renglones cuya orden NO tiene ninguna entrega vinculada (sin dato de fecha
# real) quedan fuera de N: no se puede evaluar si llegaron a tiempo, y por
# consistencia tampoco se cuentan para IF ni para el denominador de OTIF $.
# Se acumula por mes (fecha de creacion de la orden), proveedor e item, para que
# el tablero recalcule los % en el navegador segun los filtros de año/mes/
# proveedor/item ya existentes (se suman siempre los conteos brutos, nunca los %).
# ════════════════════════════════════════════════════════════════

# ── PASE 1: Historico de compras por item.xlsx -> estado on-time por Nro orden ──
print("PASE 1: Historico de compras (estado on-time por Nro orden)...")
wb2 = openpyxl.load_workbook(base + "Histórico de compras por item.xlsx", data_only=True, read_only=True)
ws2 = wb2.active

hist_orders_ontime = defaultdict(list)  # nro_orden -> [bool on_time, ...] de cada entrega vinculada

headers2 = None
n2 = 0
n_linked = 0
for i, row in enumerate(ws2.iter_rows(values_only=True)):
    if i == 0:
        headers2 = [str(c).strip() if c else '' for c in row]
        continue
    r = dict(zip(headers2, row))
    n2 += 1
    docto_orden = r.get('Docto. orden')
    if docto_orden in (None, '', 'None'):
        continue
    fecha = safe_date(r.get('Fecha'))
    f_prom = safe_date(r.get('Fecha de entrega item'))
    f_oc = safe_date(r.get('Fecha documento OC'))
    if not fecha or not f_prom or not f_oc:
        continue
    real_days = (fecha - f_oc).days
    if not (0 <= real_days <= 365):
        continue  # vinculo/fecha sospechosa (posible error de datos) -> se descarta

    hist_orders_ontime[docto_orden].append(fecha <= f_prom)
    n_linked += 1

wb2.close()
print(f"  Filas historico: {n2}  Filas vinculadas validas: {n_linked}  Ordenes con receipt: {len(hist_orders_ontime)}")

# Estado on-time por orden: True solo si TODAS sus entregas vinculadas llegaron a tiempo
order_on_time = {nro: all(flags) for nro, flags in hist_orders_ontime.items()}

# ── PASE 2: Orden de compra por item.xlsx -> N, IF, OT (por orden), OTIF y $ ──
print("PASE 2: Orden de compra por item (poblacion unificada N)...")
wb1 = openpyxl.load_workbook(base + "Orden de compra por item.xlsx", data_only=True, read_only=True)
ws1 = wb1.active

overall = new_bucket()
byProv = defaultdict(new_bucket)
byItem = defaultdict(new_bucket)

headers1 = None
n1 = 0
n_anulado = 0
n_sin_receipt = 0
for i, row in enumerate(ws1.iter_rows(values_only=True)):
    if i == 0:
        headers1 = [str(c).strip() if c else '' for c in row]
        continue
    r = dict(zip(headers1, row))
    n1 += 1
    nro = r.get('Nro orden')
    estado = str(r.get('Estado') or '').strip()
    if estado == 'Anulado':
        n_anulado += 1
        continue
    if nro is None:
        continue

    ot = order_on_time.get(nro)
    if ot is None:
        n_sin_receipt += 1
        continue  # sin entrega vinculada -> no se puede evaluar OT, se excluye de N

    cant_ord = safe_float(r.get('Cant. ordenada'))
    cant_ent = safe_float(r.get('Cant. entrada'))
    valor = safe_float(r.get('Valor neto'))
    prov = str(r.get('Razón social proveedor') or r.get('Razon social proveedor') or '').strip()
    item = str(r.get('Desc. item') or '').strip()
    f_orden = safe_date(r.get('Fecha orden'))
    if not f_orden:
        continue
    key = f"{f_orden.year}-{f_orden.month:02d}"
    if key not in month_idx:
        continue
    mi = month_idx[key]

    if_ok = cant_ord > 0 and cant_ent >= cant_ord - 1e-6
    otif_ok = if_ok and ot

    targets = [overall]
    if prov: targets.append(byProv[prov])
    if item: targets.append(byItem[item])
    for b in targets:
        b["n"][mi] += 1
        b["valorEval"][mi] += valor
        if if_ok: b["ifOk"][mi] += 1
        if ot: b["otOk"][mi] += 1
        if otif_ok:
            b["otifOk"][mi] += 1
            b["valorOtifOk"][mi] += valor

wb1.close()
print(f"  Filas OC: {n1}  Anuladas (excluidas): {n_anulado}  Sin receipt vinculado (excluidas de N): {n_sin_receipt}")


def round_bucket(b):
    return {k: [round(v, 4) for v in arr] for k, arr in b.items()}


byProv_r = {p: round_bucket(b) for p, b in byProv.items() if p}
byItem_r = {it: round_bucket(b) for it, b in byItem.items() if it}
overall_r = round_bucket(overall)


def tot(arr): return sum(arr)


n, ifOk, otOk, otifOk = tot(overall_r["n"]), tot(overall_r["ifOk"]), tot(overall_r["otOk"]), tot(overall_r["otifOk"])
vEval, vOk = tot(overall_r["valorEval"]), tot(overall_r["valorOtifOk"])

print(f"\n  N (denominador comun, lineas evaluables): {int(n)}")
print(f"  OT %:   {otOk/n*100:.1f}%  ({int(otOk)}/{int(n)})")
print(f"  IF %:   {ifOk/n*100:.1f}%  ({int(ifOk)}/{int(n)})")
print(f"  OTIF %: {otifOk/n*100:.1f}%  ({int(otifOk)}/{int(n)})")
print(f"  OTIF $ %: {vOk/vEval*100:.1f}%   valor OTIF: {vOk:,.0f} / valor evaluable: {vEval:,.0f}")
print(f"  Proveedores con datos: {len(byProv_r)}   Items con datos: {len(byItem_r)}")

OTIF_DATA = {
    "monthly": overall_r,
    "byProv": byProv_r,
    "byItem": byItem_r,
}

with open(base + "extract/otif_data.json", "w", encoding="utf-8") as f:
    json.dump(OTIF_DATA, f, ensure_ascii=False, separators=(',', ':'))

print("\nGuardado otif_data.json")
