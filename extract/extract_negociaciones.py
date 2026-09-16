import openpyxl, json

BASE = "C:/Users/Montolivo/Dropbox/Asesorias PYMES/1. Proyectos Actuales/2. MONTOLIVO/2. Proyecto/Indicadores Montolivo/Compras/Bases para actualizar Claude/"

def safe_num(v):
    return float(v) if isinstance(v, (int, float)) else None

def extract_resumen(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb['RESUMEN'] if 'RESUMEN' in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    kpi_row = rows[2]  # fila 3 (0-indexed 2)
    kpis = {
        "ahorroMensualTotal": safe_num(kpi_row[0]),
        "proyeccionAnual": safe_num(kpi_row[3]),
        "pctAhorroSobreCompraBase": safe_num(kpi_row[6]),
    }

    productos = []
    for row in rows[6:]:  # a partir de fila 7 (0-indexed 6)
        prod = row[0]
        if prod is None or (isinstance(prod, str) and prod.strip().upper() == 'TOTAL'):
            break
        if isinstance(prod, (int, float)) and prod == 0:
            break
        estado_raw = row[6] if len(row) > 6 else None
        estado = str(estado_raw).strip() if estado_raw not in (None, 0, '0') else 'SIN ESTADO'
        productos.append({
            "producto": str(prod).strip(),
            "precioActual": safe_num(row[1]),
            "precioNuevo": safe_num(row[2]),
            "ahorroMensual": safe_num(row[3]),
            "reduccionPct": safe_num(row[4]),
            "ahorroAnual": safe_num(row[5]),
            "estado": estado,
        })
    return {"kpis": kpis, "productos": productos}

out = extract_resumen(BASE + "AHORROS A SEPT.xlsx")
print(f"Productos: {len(out['productos'])}  KPIs: {out['kpis']}")

with open(BASE + "extract/negociaciones_seed.json", "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, separators=(',', ':'))
print("Guardado negociaciones_seed.json")
